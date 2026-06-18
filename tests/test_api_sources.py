"""
Phase A behavioral guarantees for /sources endpoints.

Guarantees under test (ROADMAP.md feat/api-sources-register):
  1. POST /sources with a valid CSV returns a source record (ok: true).
  2. POST /sources with a valid .xlsx returns a source record (ok: true).
  3. POST /sources with an unsupported file type returns 422, not a 500.
  4. POST /sources when create_source fails returns a failure payload, not a 500.
  5. GET /sources returns the list of registered sources.
"""
import io
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from pipeui.api.sources import router, get_conn


@pytest.fixture
def client(db):
    """TestClient wired to the test's in-memory DuckDB sandbox."""
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_conn] = lambda: db
    yield TestClient(app)


def _csv_file(content: str = "id,name,value\n1,alice,10\n2,bob,20\n"):
    return ("test.csv", io.BytesIO(content.encode()), "text/csv")


def _xlsx_file():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "name", "value"])
    ws.append([1, "alice", 10])
    ws.append([2, "bob", 20])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return ("test.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@pytest.mark.integration
def test_post_sources_csv_returns_source_record(client):
    """POST /sources with a valid CSV returns ok:true and a source record."""
    name, data, mime = _csv_file()
    resp = client.post(
        "/sources",
        data={"source_name": "test_csv", "primary_key": "id"},
        files={"file": (name, data, mime)},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["ok"] is True
    assert body["source"]["source_name"] == "test_csv"
    assert body["source"]["primary_key"] == "id"
    assert isinstance(body["source"]["columns"], list)


@pytest.mark.integration
def test_post_sources_xlsx_returns_source_record(client):
    """POST /sources with a valid .xlsx returns ok:true and a source record."""
    name, data, mime = _xlsx_file()
    resp = client.post(
        "/sources",
        data={"source_name": "test_xlsx", "primary_key": "id"},
        files={"file": (name, data, mime)},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["ok"] is True
    assert body["source"]["source_name"] == "test_xlsx"


@pytest.mark.integration
def test_post_sources_bad_file_type_returns_422_not_500(client):
    """POST /sources with an unsupported file type returns 422, never 500."""
    resp = client.post(
        "/sources",
        data={"source_name": "bad", "primary_key": "id"},
        files={"file": ("bad.json", io.BytesIO(b"{}"), "application/json")},
    )
    assert resp.status_code == 422
    assert resp.status_code != 500


@pytest.mark.integration
def test_post_sources_failure_returns_failure_payload_not_500(client):
    """POST /sources when create_source fails returns a failure payload, not a 500."""
    name, data, mime = _csv_file()
    resp = client.post(
        "/sources",
        data={"source_name": "fail_test", "primary_key": "id", "ingestion_method": "INVALID"},
        files={"file": (name, data, mime)},
    )
    assert resp.status_code == 422
    body = resp.json()
    assert body["ok"] is False
    assert isinstance(body["errors"], list)
    assert len(body["errors"]) > 0


@pytest.mark.integration
def test_get_sources_returns_registered_sources(client):
    """GET /sources returns all registered sources."""
    for i in range(2):
        name, data, mime = _csv_file(f"id,val\n{i},x\n")
        client.post(
            "/sources",
            data={"source_name": f"src_{i}", "primary_key": "id"},
            files={"file": (name, data, mime)},
        )

    resp = client.get("/sources")
    assert resp.status_code == 200
    sources = resp.json()
    assert len(sources) >= 2
    names = [s["source_name"] for s in sources]
    assert "src_0" in names
    assert "src_1" in names


@pytest.mark.integration
def test_get_source_detail_distinct_pk_count_before_and_after_ingestion(client, db):
    """GET /sources/{id} includes distinct_pk_count: null before ingest, correct count after.

    Guarantee: distinct_pk_count is null when instance table does not yet exist,
    and equals the number of distinct PK values after ingestion.
    When a source has duplicate PKs inserted directly (bypassing the PK constraint),
    distinct_pk_count < row_count.
    """
    # Register a source
    name, data, mime = _csv_file("id,name\n1,alice\n2,bob\n3,carol\n")
    resp = client.post(
        "/sources",
        data={"source_name": "pk_distinct_test", "primary_key": "id"},
        files={"file": (name, data, mime)},
    )
    assert resp.status_code == 200
    source_id = resp.json()["source"]["source_id"]

    # Before ingestion: distinct_pk_count is null (instance table does not exist)
    resp = client.get(f"/sources/{source_id}")
    assert resp.status_code == 200
    detail = resp.json()
    assert "distinct_pk_count" in detail
    assert detail["distinct_pk_count"] is None

    # Ingest 3 rows with unique PKs
    ingest_file = ("data.csv", io.BytesIO(b"id,name\n1,alice\n2,bob\n3,carol\n"), "text/csv")
    resp = client.post(f"/sources/{source_id}/ingest", files={"file": ingest_file})
    assert resp.status_code == 200

    # After ingestion: distinct_pk_count equals row_count when all PKs are unique
    resp = client.get(f"/sources/{source_id}")
    assert resp.status_code == 200
    detail = resp.json()
    assert detail["distinct_pk_count"] == 3
    assert detail["row_count"] == 3
    assert detail["distinct_pk_count"] == detail["row_count"]

    # Simulate duplicate PKs by recreating the instance table without the PK
    # constraint and inserting a duplicate row — this tests the scenario that the
    # warning badge detects (distinct_pk_count < row_count).
    from pipeui.backend.data.base.tables import instance_table_name
    import uuid
    tname = instance_table_name(uuid.UUID(source_id))
    # Recreate without PK constraint so duplicates can be inserted
    db.execute(f'CREATE TABLE "{tname}_nokey" AS SELECT * FROM "{tname}"')
    db.execute(f'INSERT INTO "{tname}_nokey" VALUES (1, \'alice_dup\')')
    db.execute(f'DROP TABLE "{tname}"')
    db.execute(f'ALTER TABLE "{tname}_nokey" RENAME TO "{tname}"')

    resp = client.get(f"/sources/{source_id}")
    assert resp.status_code == 200
    detail = resp.json()
    assert detail["row_count"] == 4
    assert detail["distinct_pk_count"] == 3
    assert detail["row_count"] > detail["distinct_pk_count"]


# ---------------------------------------------------------------------------
# POST /sources/peek-columns — header-only column read (no client-side parsing)
# ---------------------------------------------------------------------------

@pytest.mark.integration
def test_peek_columns_csv_returns_headers(client):
    """peek-columns returns the CSV header row's column names."""
    name, data, mime = _csv_file()
    resp = client.post("/sources/peek-columns", files={"file": (name, data, mime)})
    assert resp.status_code == 200
    assert resp.json()["columns"] == ["id", "name", "value"]


@pytest.mark.integration
def test_peek_columns_xlsx_returns_headers(client):
    """peek-columns returns the XLSX header row's column names (openpyxl read-only)."""
    name, data, mime = _xlsx_file()
    resp = client.post("/sources/peek-columns", files={"file": (name, data, mime)})
    assert resp.status_code == 200
    assert resp.json()["columns"] == ["id", "name", "value"]


@pytest.mark.integration
def test_peek_columns_unsupported_type_returns_422(client):
    """An unsupported extension returns 422, not 500."""
    resp = client.post(
        "/sources/peek-columns",
        files={"file": ("notes.txt", io.BytesIO(b"id,name\n1,a\n"), "text/plain")},
    )
    assert resp.status_code == 422


def test_peek_header_columns_helper_reads_header_and_drops_blanks(tmp_path):
    """The helper returns only the first row and drops blank/whitespace header cells,
    for both CSV and XLSX — without materializing the rest of the file."""
    from pipeui.backend.domain.sources.create import peek_header_columns

    csv_path = tmp_path / "h.csv"
    csv_path.write_text("id,amount, ,region\n1,10,x,west\n2,20,y,east\n")
    assert peek_header_columns(str(csv_path)) == ["id", "amount", "region"]

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "amount", None, "region"])
    for i in range(5000):  # many rows; only the header should be read back
        ws.append([i, i * 10, "x", "west"])
    xlsx_path = tmp_path / "h.xlsx"
    wb.save(str(xlsx_path))
    assert peek_header_columns(str(xlsx_path)) == ["id", "amount", "region"]
